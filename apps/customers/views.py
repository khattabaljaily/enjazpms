import csv
import io
import json
from decimal import Decimal

from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.db.models import DecimalField, Exists, OuterRef, Q, Sum, Value
from django.db.models.functions import Coalesce
from django.http import HttpResponse, HttpResponseNotAllowed, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.http import require_POST

from .forms import CustomerForm
from .models import Customer
from apps.sales.models import CustomerLedger
from apps.sales.services import _apply_customer_ledger
from apps.treasury.models import Treasury, TreasuryMovement
from apps.treasury.services import post_treasury_disbursement, post_treasury_receipt


def _ensure_tenant(request):
    tenant = getattr(request, 'tenant', None)
    if not tenant:
        return None
    return tenant


@login_required
def customer_list(request):
    tenant = _ensure_tenant(request)
    if not tenant:
        return redirect('core:no_tenant')

    qs = Customer.objects.for_tenant(tenant)
    total = qs.count()
    active = qs.filter(is_active=True).count()
    inactive = total - active

    context = {
        'form': CustomerForm(),
        'stats': {
            'total': total,
            'active': active,
            'inactive': inactive,
        },
    }
    return render(request, 'customers/customer_list.html', context)


def _serialize_form_errors(form):
    return {field: [str(error) for error in errors] for field, errors in form.errors.items()}


def _json_error(message, status=400):
    return JsonResponse({'success': False, 'message': message}, status=status)


def _json_ok(data=None, msg='تمت العملية بنجاح'):
    payload = {'success': True, 'message': msg}
    if data is not None:
        payload['data'] = data
    return JsonResponse(payload)


@login_required
def customer_table_api(request):
    tenant = _ensure_tenant(request)
    if not tenant:
        return JsonResponse({'success': False, 'message': 'لا يوجد نشاط تجاري'}, status=400)

    draw = int(request.GET.get('draw', 1))
    start = int(request.GET.get('start', 0))
    length = int(request.GET.get('length', 25))
    search_value = request.GET.get('search[value]', '').strip()
    status = request.GET.get('status', '').strip()

    queryset = Customer.objects.for_tenant(tenant).annotate(
        ledger_total=Coalesce(
            Sum('ledger_entries__amount'),
            Value(0),
            output_field=DecimalField(max_digits=14, decimal_places=2),
        )
    )
    records_total = queryset.count()

    if status == 'active':
        queryset = queryset.filter(is_active=True)
    elif status == 'inactive':
        queryset = queryset.filter(is_active=False)

    if search_value:
        queryset = queryset.filter(
            Q(name__icontains=search_value)
            | Q(code__icontains=search_value)
            | Q(phone__icontains=search_value)
            | Q(email__icontains=search_value)
            | Q(city__icontains=search_value)
        )

    records_filtered = queryset.count()

    order_column_index = request.GET.get('order[0][column]', '0')
    order_dir = request.GET.get('order[0][dir]', 'asc')
    order_column_name = request.GET.get(f'columns[{order_column_index}][data]', 'created_at')

    allowed_order_fields = {
        'code': 'code',
        'name': 'name',
        'phone': 'phone',
        'city': 'city',
        'opening_balance': 'opening_balance',
        'current_balance': 'opening_balance',
        'is_active': 'is_active',
        'created_at': 'created_at',
    }
    order_field = allowed_order_fields.get(order_column_name, 'created_at')
    if order_dir == 'desc':
        order_field = f'-{order_field}'

    queryset = queryset.order_by(order_field)[start:start + length]

    data = [
        {
            'id': customer.id,
            'code': customer.code,
            'name': customer.name,
            'phone': customer.phone or '-',
            'city': customer.city or '-',
            'opening_balance': str(customer.opening_balance),
            'current_balance': str((customer.opening_balance or 0) + (customer.ledger_total or 0)),
            'is_active': customer.is_active,
        }
        for customer in queryset
    ]

    return JsonResponse(
        {
            'draw': draw,
            'recordsTotal': records_total,
            'recordsFiltered': records_filtered,
            'data': data,
        }
    )


@login_required
def customer_create_api(request):
    tenant = _ensure_tenant(request)
    if not tenant:
        return JsonResponse({'success': False, 'message': 'لا يوجد نشاط تجاري'}, status=400)

    if request.method != 'POST':
        return HttpResponseNotAllowed(['POST'])

    form = CustomerForm(request.POST)
    if form.is_valid():
        customer = form.save(commit=False)
        customer.tenant = tenant
        customer.created_by = request.user
        customer.updated_by = request.user
        customer.save()
        return JsonResponse({
            'success': True,
            'message': 'تم إضافة العميل بنجاح',
            'id': customer.id,
        })

    return JsonResponse({
        'success': False,
        'message': 'يرجى التحقق من الحقول المطلوبة',
        'errors': _serialize_form_errors(form),
    }, status=400)


@login_required
def customer_detail_api(request, pk):
    tenant = _ensure_tenant(request)
    if not tenant:
        return JsonResponse({'success': False, 'message': 'لا يوجد نشاط تجاري'}, status=400)

    customer = get_object_or_404(Customer.objects.for_tenant(tenant), pk=pk)
    ledger_total = (
        CustomerLedger.objects
        .for_tenant(tenant)
        .filter(customer=customer)
        .aggregate(s=Sum('amount'))['s']
        or 0
    )
    current_balance = (customer.opening_balance or 0) + ledger_total

    return JsonResponse({
        'success': True,
        'data': {
            'id': customer.id,
            'name': customer.name,
            'phone': customer.phone,
            'email': customer.email,
            'city': customer.city,
            'address': customer.address,
            'opening_balance': str(customer.opening_balance),
            'current_balance': str(current_balance),
            'credit_limit': str(customer.credit_limit),
            'notes': customer.notes,
            'is_active': customer.is_active,
        }
    })


@login_required
def customer_transactions_api(request, pk):
    tenant = _ensure_tenant(request)
    if not tenant:
        return JsonResponse({'success': False, 'message': 'لا يوجد نشاط تجاري'}, status=400)

    customer = get_object_or_404(Customer.objects.for_tenant(tenant), pk=pk)
    entries = (
        CustomerLedger.objects
        .for_tenant(tenant)
        .filter(customer=customer)
        .order_by('-entry_date', '-created_at')[:100]
    )

    type_labels = {
        'opening': 'رصيد افتتاحي',
        'invoice': 'فاتورة آجل',
        'payment': 'سداد عميل',
        'return': 'مرتجع',
        'adjustment': 'تعديل',
    }

    data = [
        {
            'entry_date': e.entry_date.strftime('%Y-%m-%d'),
            'entry_type': e.entry_type,
            'entry_type_label': type_labels.get(e.entry_type, e.entry_type),
            'amount': str(e.amount),
            'running_balance': str(e.running_balance),
            'notes': e.notes or '—',
            'reference_type': e.reference_type or '',
            'reference_id': e.reference_id,
        }
        for e in entries
    ]

    return JsonResponse({'success': True, 'data': data})


@login_required
def customer_payments(request):
    tenant = _ensure_tenant(request)
    if not tenant:
        return redirect('core:no_tenant')

    customers = Customer.objects.for_tenant(tenant).filter(is_active=True).annotate(
        ledger_total=Coalesce(
            Sum('ledger_entries__amount', output_field=DecimalField(max_digits=14, decimal_places=2)),
            Value(0, output_field=DecimalField(max_digits=14, decimal_places=2)),
            output_field=DecimalField(max_digits=14, decimal_places=2),
        )
    ).order_by('name')
    treasuries = Treasury.objects.for_tenant(tenant).filter(is_active=True).order_by('name')
    stats = CustomerLedger.objects.for_tenant(tenant).filter(entry_type='payment').aggregate(
        total=Coalesce(
            Sum('amount', output_field=DecimalField(max_digits=14, decimal_places=2)),
            Value(0, output_field=DecimalField(max_digits=14, decimal_places=2)),
            output_field=DecimalField(max_digits=14, decimal_places=2),
        ),
        cash=Coalesce(
            Sum('amount', filter=Q(reference_type='customer_payment_cash'), output_field=DecimalField(max_digits=14, decimal_places=2)),
            Value(0, output_field=DecimalField(max_digits=14, decimal_places=2)),
            output_field=DecimalField(max_digits=14, decimal_places=2),
        ),
        bank=Coalesce(
            Sum('amount', filter=Q(reference_type='customer_payment_bank'), output_field=DecimalField(max_digits=14, decimal_places=2)),
            Value(0, output_field=DecimalField(max_digits=14, decimal_places=2)),
            output_field=DecimalField(max_digits=14, decimal_places=2),
        ),
    )

    def positive(value):
        return abs(value) if value is not None else 0

    context = {
        'customers': customers,
        'treasuries': treasuries,
        'stats': {
            'total': CustomerLedger.objects.for_tenant(tenant).filter(entry_type='payment').count(),
            'total_amount': positive(stats['total']),
            'cash_amount': positive(stats['cash']),
            'bank_amount': positive(stats['bank']),
        },
        'today': timezone.now().date().isoformat(),
    }
    return render(request, 'customers/payment_list.html', context)


@login_required
def customer_payments_table_api(request):
    tenant = _ensure_tenant(request)
    if not tenant:
        return _json_error('لا يوجد نشاط تجاري')

    draw = int(request.GET.get('draw', 1))
    start = int(request.GET.get('start', 0))
    length = int(request.GET.get('length', 25))
    search_value = request.GET.get('search[value]', '').strip()
    customer_filter = request.GET.get('customer_id', '')
    method_filter = request.GET.get('payment_method', '')

    qs = CustomerLedger.objects.for_tenant(tenant).filter(entry_type='payment')
    total = qs.count()

    if customer_filter:
        qs = qs.filter(customer_id=customer_filter)
    if method_filter:
        qs = qs.filter(reference_type=f'customer_payment_{method_filter}')

    if search_value:
        qs = qs.filter(
            Q(customer__name__icontains=search_value)
            | Q(notes__icontains=search_value)
            | Q(reference_type__icontains=search_value)
        )

    filtered_total = qs.count()

    order_col = request.GET.get('order[0][column]', None)
    order_dir = request.GET.get('order[0][dir]', 'desc')
    col_map = {
        '0': 'entry_date',
        '1': 'customer__name',
        '2': 'amount',
        '3': 'reference_type',
    }
    if order_col is None:
        order_field = '-id'
    else:
        order_field = col_map.get(order_col, 'id')
        if order_dir == 'desc':
            order_field = f'-{order_field}'
    qs = qs.order_by(order_field)

    cancel_qs = CustomerLedger.objects.for_tenant(tenant).filter(
        reference_type='customer_payment_cancel',
        reference_id=OuterRef('pk'),
    )
    qs = qs.annotate(is_canceled=Exists(cancel_qs))

    cancel_qs = CustomerLedger.objects.for_tenant(tenant).filter(
        reference_type='customer_payment_cancel',
        reference_id=OuterRef('pk'),
    )
    qs = qs.annotate(is_canceled=Exists(cancel_qs))

    page_qs = list(qs[start: start + length])
    sale_payment_ids = [entry.reference_id for entry in page_qs
                        if entry.reference_type == 'sale_payment' and entry.reference_id]
    sale_payment_methods = {}
    if sale_payment_ids:
        from apps.sales.models import SalePayment
        sale_payment_methods = {
            p.id: p.payment_method
            for p in SalePayment.objects.for_tenant(tenant).filter(id__in=sale_payment_ids)
        }

    data = []
    for entry in page_qs:
        if entry.reference_type == 'customer_payment_cash':
            method_label = 'نقداً'
        elif entry.reference_type == 'customer_payment_bank':
            method_label = 'بنكي'
        elif entry.reference_type == 'sale_payment':
            payment_method = sale_payment_methods.get(entry.reference_id)
            if payment_method == 'cash':
                method_label = 'نقداً'
            elif payment_method == 'bank':
                method_label = 'بنكي'
            else:
                method_label = '—'
        else:
            method_label = '—'

        if entry.is_canceled:
            method_label += ' — ملغاة'
        data.append({
            'id': entry.id,
            'entry_date': entry.entry_date.strftime('%Y-%m-%d'),
            'customer': entry.customer.name,
            'amount': str(entry.amount),
            'payment_method': method_label,
            'notes': entry.notes or '—',
            'entry_type': entry.entry_type,
            'is_canceled': bool(entry.is_canceled),
        })

    return JsonResponse({
        'draw': draw,
        'recordsTotal': total,
        'recordsFiltered': filtered_total,
        'data': data,
    })


@login_required
def customer_payment_detail_api(request, pk):
    tenant = _ensure_tenant(request)
    if not tenant:
        return _json_error('لا يوجد نشاط تجاري')

    payment = get_object_or_404(
        CustomerLedger.objects.for_tenant(tenant).select_related('customer'),
        entry_type='payment',
        pk=pk,
    )

    cancellation = CustomerLedger.objects.for_tenant(tenant).filter(
        reference_type='customer_payment_cancel',
        reference_id=payment.id,
    ).first()
    cash_treasury = None
    if payment.reference_type == 'customer_payment_cash':
        treasury_movement = TreasuryMovement.objects.for_tenant(tenant).filter(
            reference_type='customer_payment_cash',
            reference_id=payment.id,
        ).select_related('treasury').first()
        if treasury_movement:
            cash_treasury = treasury_movement.treasury.name

    response_data = {
        'id': payment.id,
        'entry_date': payment.entry_date.strftime('%Y-%m-%d'),
        'customer': payment.customer.name,
        'amount': str(payment.amount),
        'payment_method': 'نقداً' if payment.reference_type == 'customer_payment_cash' else 'بنكي',
        'notes': payment.notes or '—',
        'is_canceled': bool(cancellation),
        'cancellation_note': cancellation.notes if cancellation else '',
        'cancellation_date': cancellation.entry_date.strftime('%Y-%m-%d') if cancellation else None,
        'cash_treasury': cash_treasury,
    }
    return _json_ok(data=response_data)


@login_required
@require_POST
def customer_payment_create_api(request):
    tenant = _ensure_tenant(request)
    if not tenant:
        return _json_error('لا يوجد نشاط تجاري')

    if request.content_type != 'application/json':
        return _json_error('بيانات غير صالحة', status=400)

    try:
        body = json.loads(request.body)
        customer_id = int(body.get('customer_id'))
        amount = Decimal(str(body.get('amount')))
        payment_date = body.get('payment_date') or timezone.now().date().isoformat()
        method = body.get('method', 'cash')
        treasury_id = body.get('treasury_id')
        reference = str(body.get('reference', '') or '').strip()
        notes = str(body.get('notes', '') or '').strip()
    except (TypeError, ValueError, json.JSONDecodeError) as e:
        return _json_error(f'بيانات الدفعة غير صالحة: {e}')

    if amount <= 0:
        return _json_error('المبلغ يجب أن يكون أكبر من الصفر')

    customer = get_object_or_404(Customer.objects.for_tenant(tenant), pk=customer_id)
    reference_type = 'customer_payment_bank' if method == 'bank' else 'customer_payment_cash'
    note_text = notes
    if reference:
        note_text = f"{note_text} | مرجع: {reference}" if note_text else f"مرجع: {reference}"
    if not note_text:
        note_text = 'سداد عميل'

    try:
        with transaction.atomic():
            payment_entry = _apply_customer_ledger(
                tenant=tenant,
                customer=customer,
                amount=-amount,
                entry_type='payment',
                reference_type=reference_type,
                reference_id=None,
                date=payment_date,
                notes=note_text,
            )

            if method == 'cash':
                if not treasury_id:
                    raise ValueError('يجب اختيار الخزينة عند دفع نقداً')
                treasury = get_object_or_404(Treasury.objects.for_tenant(tenant), pk=int(treasury_id))
                movement = post_treasury_receipt(
                    tenant=tenant,
                    amount=amount,
                    date=payment_date,
                    reference_type='customer_payment_cash',
                    reference_id=payment_entry.id if payment_entry else None,
                    description=f'دفعة عميل {customer.name}',
                    user=request.user,
                    treasury=treasury,
                )
                if not movement:
                    raise ValueError('تعذر تسجيل حركة الخزينة')
    except ValueError as e:
        return _json_error(str(e), status=400)
    except Exception as e:
        return _json_error('تعذر تسجيل الدفعة، حاول مرة أخرى')

    balance = (
        CustomerLedger.objects.for_tenant(tenant)
        .filter(customer=customer)
        .aggregate(s=Sum('amount'))['s'] or 0
    )
    current_balance = (customer.opening_balance or 0) + balance

    return _json_ok(data={'current_balance': str(current_balance)}, msg='تم تسجيل دفعة العميل بنجاح')


@login_required
@require_POST
def customer_payment_cancel_api(request, pk):
    tenant = _ensure_tenant(request)
    if not tenant:
        return _json_error('لا يوجد نشاط تجاري')

    payment = get_object_or_404(
        CustomerLedger.objects.for_tenant(tenant).filter(entry_type='payment'),
        pk=pk,
    )
    reverse_notes = f"إلغاء دفعة عميل — {payment.notes or ''}".strip()
    with transaction.atomic():
        if payment.reference_type == 'customer_payment_cash':
            treasury_movement = TreasuryMovement.objects.for_tenant(tenant).filter(
                reference_type='customer_payment_cash',
                reference_id=payment.id,
            ).first()
            if treasury_movement:
                post_treasury_disbursement(
                    tenant=tenant,
                    amount=abs(payment.amount),
                    date=timezone.now().date(),
                    reference_type='customer_payment_cash_cancel',
                    reference_id=payment.id,
                    description=f'إلغاء دفعة عميل {payment.customer.name}',
                    user=request.user,
                    treasury=treasury_movement.treasury,
                )

        _apply_customer_ledger(
            tenant=tenant,
            customer=payment.customer,
            amount=-payment.amount,
            entry_type='adjustment',
            reference_type='customer_payment_cancel',
            reference_id=payment.id,
            date=timezone.now().date(),
            notes=reverse_notes,
        )

    return _json_ok(msg='تم إلغاء الدفعة واستعادة مديونية العميل')


@login_required
def customer_update_api(request, pk):
    tenant = _ensure_tenant(request)
    if not tenant:
        return JsonResponse({'success': False, 'message': 'لا يوجد نشاط تجاري'}, status=400)

    if request.method != 'POST':
        return HttpResponseNotAllowed(['POST'])

    customer = get_object_or_404(Customer.objects.for_tenant(tenant), pk=pk)
    form = CustomerForm(request.POST, instance=customer)

    if form.is_valid():
        customer = form.save(commit=False)
        customer.updated_by = request.user
        customer.save()
        return JsonResponse({
            'success': True,
            'message': 'تم تعديل بيانات العميل بنجاح',
        })

    return JsonResponse({
        'success': False,
        'message': 'يرجى التحقق من الحقول المطلوبة',
        'errors': _serialize_form_errors(form),
    }, status=400)


@login_required
def customer_delete_api(request, pk):
    tenant = _ensure_tenant(request)
    if not tenant:
        return JsonResponse({'success': False, 'message': 'لا يوجد نشاط تجاري'}, status=400)

    if request.method != 'POST':
        return HttpResponseNotAllowed(['POST'])

    customer = get_object_or_404(Customer.objects.for_tenant(tenant), pk=pk)
    customer.delete()
    return JsonResponse({
        'success': True,
        'message': 'تم حذف العميل بنجاح',
    })


@login_required
def customer_create(request):
    return redirect('customers:list')


@login_required
def customer_import_api(request):
    """Import customers from Excel/CSV file"""
    tenant = _ensure_tenant(request)
    if not tenant:
        return JsonResponse({'success': False, 'message': 'لا يوجد نشاط تجاري'}, status=400)

    if request.method != 'POST':
        return HttpResponseNotAllowed(['POST'])

    if 'file' not in request.FILES:
        return JsonResponse({'success': False, 'message': 'لم يتم رفع أي ملف'}, status=400)

    file = request.FILES['file']
    
    # Validate file extension
    if not file.name.endswith(('.csv', '.xlsx', '.xls')):
        return JsonResponse({'success': False, 'message': 'نوع الملف غير مدعوم'}, status=400)

    try:
        imported_count = 0
        errors = []

        if file.name.endswith('.csv'):
            # Handle CSV
            decoded_file = file.read().decode('utf-8-sig')
            csv_reader = csv.DictReader(io.StringIO(decoded_file))
            
            for row_num, row in enumerate(csv_reader, start=2):
                try:
                    Customer.objects.create(
                        tenant=tenant,
                        name=row.get('name', '').strip() or row.get('الاسم', '').strip(),
                        phone=row.get('phone', '').strip() or row.get('الهاتف', '').strip() or None,
                        email=row.get('email', '').strip() or row.get('البريد', '').strip() or None,
                        city=row.get('city', '').strip() or row.get('المدينة', '').strip() or None,
                        address=row.get('address', '').strip() or row.get('العنوان', '').strip() or None,
                        opening_balance=float(row.get('opening_balance', 0) or row.get('الرصيد', 0) or 0),
                        credit_limit=float(row.get('credit_limit', 0) or row.get('حد_الائتمان', 0) or 0),
                        is_active=True
                    )
                    imported_count += 1
                except Exception as e:
                    errors.append(f'الصف {row_num}: {str(e)}')
        else:
            # Handle Excel - requires openpyxl
            try:
                import openpyxl
            except ImportError:
                return JsonResponse({
                    'success': False, 
                    'message': 'مكتبة openpyxl غير مثبتة. الرجاء تثبيتها أولاً'
                }, status=500)

            wb = openpyxl.load_workbook(file)
            ws = wb.active
            
            # Get headers from first row
            headers = [cell.value for cell in ws[1]]
            
            for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                try:
                    data = dict(zip(headers, row))
                    Customer.objects.create(
                        tenant=tenant,
                        name=str(data.get('name', '') or data.get('الاسم', '')).strip(),
                        phone=str(data.get('phone', '') or data.get('الهاتف', '')).strip() or None,
                        email=str(data.get('email', '') or data.get('البريد', '')).strip() or None,
                        city=str(data.get('city', '') or data.get('المدينة', '')).strip() or None,
                        address=str(data.get('address', '') or data.get('العنوان', '')).strip() or None,
                        opening_balance=float(data.get('opening_balance', 0) or data.get('الرصيد', 0) or 0),
                        credit_limit=float(data.get('credit_limit', 0) or data.get('حد_الائتمان', 0) or 0),
                        is_active=True
                    )
                    imported_count += 1
                except Exception as e:
                    errors.append(f'الصف {row_num}: {str(e)}')

        message = f'تم استيراد {imported_count} عميل بنجاح'
        if errors:
            message += f'. حدثت {len(errors)} أخطاء'

        return JsonResponse({
            'success': True,
            'message': message,
            'imported': imported_count,
            'errors': errors[:10]  # Return first 10 errors only
        })

    except Exception as e:
        return JsonResponse({
            'success': False,
            'message': f'حدث خطأ أثناء الاستيراد: {str(e)}'
        }, status=500)


@login_required
def customer_export_api(request):
    """Export customers to CSV file"""
    tenant = _ensure_tenant(request)
    if not tenant:
        return JsonResponse({'success': False, 'message': 'لا يوجد نشاط تجاري'}, status=400)

    # Create CSV response
    response = HttpResponse(content_type='text/csv; charset=utf-8-sig')
    response['Content-Disposition'] = 'attachment; filename="customers.csv"'
    
    # Add BOM for Excel UTF-8 support
    response.write('\ufeff')
    
    writer = csv.writer(response)
    
    # Write headers
    writer.writerow([
        'الاسم', 'الكود', 'الهاتف', 'البريد', 'المدينة', 'العنوان',
        'الرصيد الافتتاحي', 'حد الائتمان', 'الملاحظات', 'نشط'
    ])
    
    # Write data
    customers = Customer.objects.for_tenant(tenant).order_by('name')
    for customer in customers:
        writer.writerow([
            customer.name,
            customer.code,
            customer.phone or '',
            customer.email or '',
            customer.city or '',
            customer.address or '',
            customer.opening_balance,
            customer.credit_limit,
            customer.notes or '',
            'نعم' if customer.is_active else 'لا'
        ])
    
    return response


@login_required
def download_template(request):
    """Download CSV template for import"""
    response = HttpResponse(content_type='text/csv; charset=utf-8-sig')
    response['Content-Disposition'] = 'attachment; filename="customers_template.csv"'
    
    # Add BOM for Excel UTF-8 support
    response.write('\ufeff')
    
    writer = csv.writer(response)
    
    # Write headers
    writer.writerow([
        'الاسم', 'الهاتف', 'البريد', 'المدينة', 'العنوان',
        'الرصيد الافتتاحي', 'حد الائتمان'
    ])
    
    # Write example row
    writer.writerow([
        'أحمد محمد', '0512345678', 'ahmad@example.com', 'الرياض', 'شارع الملك فهد',
        '0', '5000'
    ])
    
    return response
