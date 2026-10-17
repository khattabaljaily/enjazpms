"""
صفحة «استيراد البيانات» الموحّدة: التصنيفات، المنتجات، العملاء، الموردون، المناديب.
شاشات القوائم نفسها لا تحمل أي زر/نافذة استيراد (التصدير فقط).
"""
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse

from apps.core.test_utils import TenantTestCase
from apps.items.models import Category


class HubCardsTests(TenantTestCase):
    subscription_plan = 'pro'

    def _hub(self):
        return self.client.get(reverse('data_import:hub'))

    def test_hub_lists_all_five_in_order(self):
        self.tenant.capabilities.has_agents_module = True
        self.tenant.capabilities.save()
        html = self._hub().content.decode()
        order = [html.find(t) for t in (
            'استيراد التصنيفات', 'استيراد المنتجات', 'استيراد العملاء', 'استيراد الموردين', 'استيراد المناديب')]
        self.assertTrue(all(i > 0 for i in order), order)
        self.assertEqual(order, sorted(order))

    def test_agents_card_hidden_without_agents_module(self):
        self.tenant.capabilities.has_agents_module = False
        self.tenant.capabilities.save()
        self.assertNotContains(self._hub(), 'استيراد المناديب')
        self.assertNotEqual(self.client.get(reverse('data_import:agent_page')).status_code, 200)


class BasicPlanAgentsTests(TenantTestCase):
    subscription_plan = 'basic'

    def test_agents_hidden_on_basic_plan(self):
        self.assertNotContains(self.client.get(reverse('data_import:hub')), 'استيراد المناديب')
        self.assertNotEqual(self.client.get(reverse('data_import:agent_page')).status_code, 200)


class ListScreensHaveNoImportTests(TenantTestCase):
    subscription_plan = 'pro'

    def test_category_and_agent_lists_have_no_import_ui(self):
        self.tenant.capabilities.has_agents_module = True
        self.tenant.capabilities.save()
        cats = self.client.get(reverse('items:category_list'))
        self.assertNotContains(cats, 'importCategoriesModal')
        self.assertNotContains(cats, 'btnOpenImportCategories')
        self.assertContains(cats, 'export_categories')
        agents = self.client.get(reverse('agents:list'))
        self.assertNotContains(agents, 'importAgentsModal')
        self.assertContains(agents, 'btnExportAgents')


class ImportCommitTests(TenantTestCase):
    subscription_plan = 'pro'

    def test_category_import_through_hub_page(self):
        page = self.client.get(reverse('data_import:category_page'))
        self.assertContains(page, reverse('items:category_import_api'))
        csv = 'الاسم,التصنيف الرئيسي\nمشروبات,\nعصائر,مشروبات\n'.encode('utf-8')
        resp = self.client.post(reverse('items:category_import_api'),
                                {'file': SimpleUploadedFile('c.csv', csv)})
        self.assertEqual(resp.json()['imported'], 2)
        self.assertEqual(Category.objects.get(tenant=self.tenant, name='عصائر').parent.name, 'مشروبات')

    def test_agent_import_through_hub_page(self):
        self.tenant.capabilities.has_agents_module = True
        self.tenant.capabilities.save()
        page = self.client.get(reverse('data_import:agent_page'))
        self.assertContains(page, reverse('agents:import_api'))
        csv = 'الاسم,الهاتف\nأحمد,0912345678\n'.encode('utf-8')
        resp = self.client.post(reverse('agents:import_api'), {'file': SimpleUploadedFile('a.csv', csv)})
        self.assertEqual(resp.json()['imported'], 1)


class SidebarPlacementTests(TenantTestCase):
    subscription_plan = 'pro'

    def test_import_link_is_top_level_item_after_settings(self):
        html = self.client.get(reverse('data_import:hub')).content.decode()
        settings_at = html.index('الإعدادات')
        import_at = html.index('استيراد البيانات', settings_at)
        # قائمة المنتجات تأتي بعد الإعدادات؛ رابط الاستيراد يجب أن يسبقها
        products_at = html.index('>\n                المنتجات', settings_at)
        self.assertLess(settings_at, import_at)
        self.assertLess(import_at, products_at)
        # قائمة الإعدادات تُغلق قبل عنصر الاستيراد: هو عنصر رئيسي مستقل لا فرعي داخلها
        self.assertLess(html.index('</ul>', settings_at), import_at)


class ImportRequiresAddPermissionTests(TenantTestCase):
    """الاستيراد يتطلب صلاحية الاستيراد وصلاحية الإضافة معًا لكل نوع بيانات."""
    subscription_plan = 'pro'

    PAGES = {
        'items': ('data_import:product_page', 'استيراد المنتجات'),
        'customers': ('data_import:customer_page', 'استيراد العملاء'),
        'suppliers': ('data_import:supplier_page', 'استيراد الموردين'),
        'agents': ('data_import:agent_page', 'استيراد المناديب'),
    }

    def _user(self, perms):
        from apps.accounts.models import PermissionGroup, User
        self.tenant.capabilities.has_agents_module = True
        self.tenant.capabilities.save()
        group = PermissionGroup.objects.create(
            tenant=self.tenant, name='g-' + '-'.join(sorted(perms)), is_active=True,
            permissions={k: True for k in perms} | {'view_items': True})
        user = User.objects.create_user(username='u-' + str(len(perms)) + str(abs(hash(tuple(sorted(perms))))),
                                        password='x12345678', tenant=self.tenant)
        user.permission_groups.add(group)
        self.client.force_login(user)
        return user

    def test_import_without_add_is_denied_and_hidden(self):
        for kind, (url_name, label) in self.PAGES.items():
            with self.subTest(kind=kind):
                self._user({f'import_{kind}'})
                self.assertNotEqual(self.client.get(reverse(url_name)).status_code, 200)
                self.assertNotContains(self.client.get(reverse('data_import:hub')), label)

    def test_add_without_import_is_denied_and_hidden(self):
        for kind, (url_name, label) in self.PAGES.items():
            with self.subTest(kind=kind):
                self._user({f'add_{kind}'})
                self.assertNotEqual(self.client.get(reverse(url_name)).status_code, 200)

    def test_import_with_add_is_allowed(self):
        for kind, (url_name, label) in self.PAGES.items():
            with self.subTest(kind=kind):
                self._user({f'import_{kind}', f'add_{kind}'})
                self.assertEqual(self.client.get(reverse(url_name)).status_code, 200)
                self.assertContains(self.client.get(reverse('data_import:hub')), label)

    def test_categories_follow_add_categories_only(self):
        self._user({'add_categories'})
        self.assertEqual(self.client.get(reverse('data_import:category_page')).status_code, 200)
        self._user({'view_categories'})
        self.assertNotEqual(self.client.get(reverse('data_import:category_page')).status_code, 200)

    def test_agent_api_also_needs_add_agents(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        self._user({'import_agents'})
        csv = 'الاسم\nأحمد\n'.encode('utf-8')
        resp = self.client.post(reverse('agents:import_api'), {'file': SimpleUploadedFile('a.csv', csv)})
        self.assertNotEqual(resp.status_code, 200)


class HardCurrencyRoundTripTests(TenantTestCase):
    """تصدير المنتجات ثم استيراد الملف نفسه في وضع العملة الصعبة يجب ألا يغيّر الأسعار."""
    subscription_plan = 'pro'

    def setUp(self):
        super().setUp()
        from decimal import Decimal
        from apps.core.test_utils import make_item
        self.tenant.hard_currency_mode = True
        self.tenant.hard_currency = 'USD'
        self.tenant.exchange_rate = Decimal('2000')
        self.tenant.save()
        Category.objects.create(tenant=self.tenant, name='عام', is_active=True)
        self.item = make_item(self.tenant, name='منتج دولاري', cost_price='3000', selling_price='4000',
                              min_selling_price='3500', cost_price_hc='1.5', selling_price_hc='2',
                              min_selling_price_hc='1.75')

    def _export(self):
        return self.client.get(reverse('items:export_api')).content

    def _import(self, raw, ai_maps_local_to_hc=False):
        from unittest import mock
        from apps.data_import import product_importer

        def fake_map(headers, schema):
            if not ai_maps_local_to_hc:
                return {}
            # ما قد يفعله الذكاء الاصطناعي: «سعر البيع» (محلي) → أقرب حقل سعر متاح = العملة الصعبة
            guess = {'سعر التكلفة': 'cost_price_hc', 'سعر البيع': 'selling_price_hc', 'أدنى سعر بيع': 'min_selling_price_hc',
                     'سعر التكلفة (عملة صعبة)': 'cost_price_hc', 'سعر البيع (عملة صعبة)': 'selling_price_hc',
                     'أدنى سعر بيع (عملة صعبة)': 'min_selling_price_hc', 'الاسم': 'name'}
            return {h: guess[h] for h in headers if h in guess}

        with mock.patch.object(product_importer, 'smart_map_headers', fake_map):
            return product_importer.import_products(self.tenant, SimpleUploadedFile('items.csv', raw), self.user)

    def _prices(self):
        self.item.refresh_from_db()
        return tuple(format(x.normalize(), 'f') for x in (
            self.item.cost_price, self.item.selling_price, self.item.min_selling_price,
            self.item.cost_price_hc, self.item.selling_price_hc, self.item.min_selling_price_hc))

    def test_round_trip_keeps_prices_when_ai_mislabels_local_columns(self):
        before = self._prices()
        self._import(self._export(), ai_maps_local_to_hc=True)
        self.assertEqual(self._prices(), before)

    def test_round_trip_keeps_prices_without_ai(self):
        before = self._prices()
        self._import(self._export())
        self.assertEqual(self._prices(), before)

    def test_local_only_file_is_treated_as_local_currency_for_new_items(self):
        csv = 'الاسم,سعر التكلفة,سعر البيع\nجديد,3000,4000\n'.encode('utf-8')
        self._import(csv, ai_maps_local_to_hc=True)
        from apps.items.models import Item
        new = Item.objects.get(tenant=self.tenant, name='جديد')
        self.assertEqual((format(new.selling_price.normalize(), 'f'), format(new.selling_price_hc.normalize(), 'f')), ('4000', '2'))

    def test_local_only_file_updates_existing_item_in_local_currency(self):
        csv = 'الاسم,سعر البيع\nمنتج دولاري,5000\n'.encode('utf-8')
        self._import(csv, ai_maps_local_to_hc=True)
        self.assertEqual(self._prices()[1], '5000')
        self.assertEqual(self._prices()[4], '2.5')
        self.assertEqual(self._prices()[0], '3000')   # التكلفة لم ترد في الملف فبقيت كما هي


class BranchProductImportTests(TenantTestCase):
    """نسخة المؤسسات: مدير الفرع يستورد المنتجات (حين تُفعَّل إضافة المنتجات للفروع) إلى فرعه فقط."""
    version_type = 'multi_branch'
    subscription_plan = 'enterprise'

    def setUp(self):
        super().setUp()
        from apps.accounts.models import User
        from apps.core.models import Branch, Settings
        from apps.core.test_utils import make_item, make_stock
        self.b1 = Branch.objects.create(tenant=self.tenant, name='فرع الشمال')
        self.b2 = Branch.objects.create(tenant=self.tenant, name='فرع الجنوب')
        self.sup = User.objects.create_user(username='bpi-sup', password='x12345678', tenant=self.tenant,
                                            branch=self.b1, is_branch_supervisor=True)
        self.stock1 = make_stock(self.tenant, name='مخزن الشمال', branch=self.b1)
        self.stock2 = make_stock(self.tenant, name='مخزن الجنوب', branch=self.b2)
        Category.objects.create(tenant=self.tenant, name='عام', is_active=True)
        self.central = make_item(self.tenant, name='مركزي', selling_price='10')
        Settings.objects.update_or_create(tenant=self.tenant, defaults={'branches_can_manage_items': True})
        self.client.force_login(self.sup)

    def _import(self, text):
        return self.client.post(reverse('data_import:product_commit'),
                                {'file': SimpleUploadedFile('p.csv', text.encode('utf-8'))}).json()

    def test_hub_and_page_available_to_branch_manager(self):
        self.assertContains(self.client.get(reverse('data_import:hub')), 'استيراد المنتجات')
        self.assertEqual(self.client.get(reverse('data_import:product_page')).status_code, 200)

    def test_disabled_flag_blocks_branch_import(self):
        from apps.core.models import Settings
        Settings.objects.filter(tenant=self.tenant).update(branches_can_manage_items=False)
        self.assertNotEqual(self.client.get(reverse('data_import:product_page')).status_code, 200)

    def test_imported_items_belong_to_the_branch(self):
        from apps.items.models import Item
        resp = self._import('الاسم,سعر البيع\nفرعي ١,5\nفرعي ٢,6\n')
        self.assertEqual(resp['created'], 2, resp)
        self.assertEqual(set(Item.objects.filter(name__startswith='فرعي').values_list('branch_id', flat=True)), {self.b1.id})

    def test_cannot_overwrite_central_item(self):
        resp = self._import('الاسم,سعر البيع\nمركزي,999\n')
        self.assertEqual(resp['created'], 0)
        self.assertTrue(resp['errors'])
        from apps.items.models import Item
        self.assertEqual(Item.objects.get(name='مركزي').selling_price, 10)

    def test_template_lists_only_branch_stocks(self):
        import io
        from openpyxl import load_workbook
        resp = self.client.get(reverse('data_import:product_template'))
        self.assertEqual(resp.status_code, 200)
        wb = load_workbook(io.BytesIO(resp.content))
        text = ' '.join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value)
        self.assertNotIn('مخزن الجنوب', text)
        self.assertIn('مخزن الشمال', text)


class HardCurrencyNoDriftTests(TenantTestCase):
    """الصعبة مخزّنة بـ4 خانات: إعادة استيراد ما صُدِّر لا يجوز أن تُحرّك السعر المحلي (10,000.00 ← 9,999.75)."""
    subscription_plan = 'pro'

    def setUp(self):
        super().setUp()
        from decimal import Decimal
        from apps.core.test_utils import make_item
        self.tenant.hard_currency_mode = True
        self.tenant.hard_currency = 'USD'
        self.tenant.exchange_rate = Decimal('3000.25')
        self.tenant.save()
        Category.objects.create(tenant=self.tenant, name='عام', is_active=True)
        rate = Decimal('3000.25')
        local = {'cost_price': Decimal('7000.00'), 'selling_price': Decimal('10000.00'),
                 'min_selling_price': Decimal('9000.00')}
        hc = {k + '_hc': (v / rate).quantize(Decimal('0.0001')) for k, v in local.items()}
        self.item = make_item(self.tenant, name='منتج كسري', **local, **hc)
        self.before = {k: getattr(self.item, k) for k in (*local, *hc)}

    def _reimport(self, raw):
        from apps.data_import.product_importer import import_products
        import_products(self.tenant, SimpleUploadedFile('p.csv', raw), self.user)
        self.item.refresh_from_db()
        return {k: getattr(self.item, k) for k in self.before}

    def test_full_export_reimport_is_exact(self):
        raw = self.client.get(reverse('items:export_api')).content
        self.assertEqual(self._reimport(raw), self.before)

    def test_hc_only_file_with_unchanged_values_is_exact(self):
        hc = self.before
        csv = ('الاسم,سعر التكلفة (عملة صعبة),سعر البيع (عملة صعبة),أدنى سعر بيع (عملة صعبة)\n'
               f"منتج كسري,{hc['cost_price_hc']},{hc['selling_price_hc']},{hc['min_selling_price_hc']}\n").encode('utf-8')
        self.assertEqual(self._reimport(csv), self.before)

    def test_changed_hc_value_still_updates_local(self):
        from decimal import Decimal
        csv = 'الاسم,سعر البيع (عملة صعبة)\nمنتج كسري,4\n'.encode('utf-8')
        after = self._reimport(csv)
        self.assertEqual(after['selling_price_hc'], Decimal('4'))
        self.assertEqual(after['selling_price'], Decimal('12001.00'))
        self.assertEqual(after['cost_price'], self.before['cost_price'])


class HybridCentralSuppliersImportTests(TenantTestCase):
    """نسخة المؤسسات بمشتريات هجينة: مدير النشاط يستورد «الموردين المركزيين» (لا يملك صلاحيات موردي الفروع)."""
    version_type = 'multi_branch'
    subscription_plan = 'enterprise'

    def setUp(self):
        super().setUp()
        from apps.core.models import Tenant
        Tenant.objects.filter(pk=self.tenant.pk).update(purchasing_mode='hybrid')
        self.tenant.refresh_from_db()
        self.client.force_login(self.user)

    def test_owner_sees_and_uses_central_suppliers_import(self):
        self.assertFalse(self.user.has_perm_key('import_suppliers'))
        self.assertContains(self.client.get(reverse('data_import:hub')), 'استيراد الموردين المركزيين')
        page = self.client.get(reverse('data_import:supplier_page'))
        self.assertContains(page, 'استيراد الموردين المركزيين')
        self.assertContains(self.client.get(reverse('core:dashboard')), 'استيراد الموردين المركزيين')
        csv = 'الاسم\nمورد مركزي أ\n'.encode('utf-8')
        resp = self.client.post(reverse('data_import:supplier_commit'), {'file': SimpleUploadedFile('s.csv', csv)})
        self.assertEqual(resp.json()['created'], 1, resp.content)
        from apps.suppliers.models import Supplier
        self.assertIsNone(Supplier.unscoped.get(tenant=self.tenant, name='مورد مركزي أ').branch_id)

    def test_branch_user_does_not_get_central_import(self):
        from apps.accounts.models import User
        from apps.core.models import Branch
        branch = Branch.objects.create(tenant=self.tenant, name='ف')
        sup = User.objects.create_user(username='hyb-sup', password='x12345678', tenant=self.tenant,
                                       branch=branch, is_branch_supervisor=True)
        self.client.force_login(sup)
        self.assertNotContains(self.client.get(reverse('data_import:hub')), 'استيراد الموردين المركزيين')

    def test_non_hybrid_owner_has_no_central_suppliers_import(self):
        from apps.core.models import Tenant
        Tenant.objects.filter(pk=self.tenant.pk).update(purchasing_mode='central')
        self.assertNotContains(self.client.get(reverse('data_import:hub')), 'استيراد الموردين المركزيين')
