from django.contrib.auth.decorators import login_required
from django.db.models import Q
from django.http import HttpResponse, HttpResponseNotAllowed, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
import csv
import io

from .forms import SupplierForm
from .models import Supplier


def _ensure_tenant(request):
    tenant = getattr(request, 'tenant', None)
    if not tenant:
        return None
    return tenant


@login_required
def supplier_list(request):
    tenant = _ensure_tenant(request)
    if not tenant:
        return redirect('core:no_tenant')

    qs = Supplier.objects.for_tenant(tenant)
    total = qs.count()
    active = qs.filter(is_active=True).count()
    inactive = total - active

    context = {
        'form': SupplierForm(),
        'stats': {
            'total': total,
            'active': active,
            'inactive': inactive,
        },
    }
    return render(request, 'suppliers/supplier_list.html', context)


def _serialize_form_errors(form):
    return {field: [str(error) for error in errors] for field, errors in form.errors.items()}


@login_required
def supplier_table_api(request):
    tenant = _ensure_tenant(request)
    if not tenant:
        return JsonResponse({'success': False, 'message': 'لا يوجد نشاط تجاري'}, status=400)

    draw = int(request.GET.get('draw', 1))
    start = int(request.GET.get('start', 0))
    length = int(request.GET.get('length', 25))
    search_value = request.GET.get('search[value]', '').strip()
    status = request.GET.get('status', '').strip()

    queryset = Supplier.objects.for_tenant(tenant)
    records_total = queryset.count()

    if status == 'active':
        queryset = queryset.filter(is_active=True)
    elif status == 'inactive':
        queryset = queryset.filter(is_active=False)

    if search_value:
        queryset = queryset.filter(
            Q(name__icontains=search_value)
            | Q(code__icontains=search_value)
            | Q(phone__icontains=search_value)
            | Q(email__icontains=search_value)
            | Q(city__icontains=search_value)
        )

    records_filtered = queryset.count()

    order_column_index = request.GET.get('order[0][column]', '0')
    order_dir = request.GET.get('order[0][dir]', 'asc')
    order_column_name = request.GET.get(f'columns[{order_column_index}][data]', 'created_at')

    allowed_order_fields = {
        'code': 'code',
        'name': 'name',
        'phone': 'phone',
        'city': 'city',
        'opening_balance': 'opening_balance',
        'is_active': 'is_active',
        'created_at': 'created_at',
    }
    order_field = allowed_order_fields.get(order_column_name, 'created_at')
    if order_dir == 'desc':
        order_field = f'-{order_field}'

    queryset = queryset.order_by(order_field)[start:start + length]

    data = [
        {
            'id': supplier.id,
            'code': supplier.code,
            'name': supplier.name,
            'phone': supplier.phone or '-',
            'city': supplier.city or '-',
            'opening_balance': str(supplier.opening_balance),
            'is_active': supplier.is_active,
        }
        for supplier in queryset
    ]

    return JsonResponse(
        {
            'draw': draw,
            'recordsTotal': records_total,
            'recordsFiltered': records_filtered,
            'data': data,
        }
    )


@login_required
def supplier_create_api(request):
    tenant = _ensure_tenant(request)
    if not tenant:
        return JsonResponse({'success': False, 'message': 'لا يوجد نشاط تجاري'}, status=400)

    if request.method != 'POST':
        return HttpResponseNotAllowed(['POST'])

    form = SupplierForm(request.POST)
    if form.is_valid():
        supplier = form.save(commit=False)
        supplier.tenant = tenant
        supplier.created_by = request.user
        supplier.updated_by = request.user
        supplier.save()
        return JsonResponse({
            'success': True,
            'message': 'تم إضافة المورد بنجاح',
            'id': supplier.id,
        })

    return JsonResponse({
        'success': False,
        'message': 'يرجى التحقق من الحقول المطلوبة',
        'errors': _serialize_form_errors(form),
    }, status=400)


@login_required
def supplier_detail_api(request, pk):
    tenant = _ensure_tenant(request)
    if not tenant:
        return JsonResponse({'success': False, 'message': 'لا يوجد نشاط تجاري'}, status=400)

    supplier = get_object_or_404(Supplier.objects.for_tenant(tenant), pk=pk)

    return JsonResponse({
        'success': True,
        'data': {
            'id': supplier.id,
            'name': supplier.name,
            'phone': supplier.phone,
            'email': supplier.email,
            'city': supplier.city,
            'address': supplier.address,
            'opening_balance': str(supplier.opening_balance),
            'credit_limit': str(supplier.credit_limit),
            'notes': supplier.notes,
            'is_active': supplier.is_active,
        }
    })


@login_required
def supplier_update_api(request, pk):
    tenant = _ensure_tenant(request)
    if not tenant:
        return JsonResponse({'success': False, 'message': 'لا يوجد نشاط تجاري'}, status=400)

    if request.method != 'POST':
        return HttpResponseNotAllowed(['POST'])

    supplier = get_object_or_404(Supplier.objects.for_tenant(tenant), pk=pk)
    form = SupplierForm(request.POST, instance=supplier)

    if form.is_valid():
        supplier = form.save(commit=False)
        supplier.updated_by = request.user
        supplier.save()
        return JsonResponse({
            'success': True,
            'message': 'تم تعديل بيانات المورد بنجاح',
        })

    return JsonResponse({
        'success': False,
        'message': 'يرجى التحقق من الحقول المطلوبة',
        'errors': _serialize_form_errors(form),
    }, status=400)


@login_required
def supplier_delete_api(request, pk):
    tenant = _ensure_tenant(request)
    if not tenant:
        return JsonResponse({'success': False, 'message': 'لا يوجد نشاط تجاري'}, status=400)

    if request.method != 'POST':
        return HttpResponseNotAllowed(['POST'])

    supplier = get_object_or_404(Supplier.objects.for_tenant(tenant), pk=pk)
    supplier.delete()
    return JsonResponse({
        'success': True,
        'message': 'تم حذف المورد بنجاح',
    })


@login_required
def supplier_create(request):
    return redirect('suppliers:list')


@login_required
def supplier_import_api(request):
    """Import suppliers from Excel/CSV file"""
    tenant = _ensure_tenant(request)
    if not tenant:
        return JsonResponse({'success': False, 'message': 'لا يوجد نشاط تجاري'}, status=400)

    if request.method != 'POST':
        return HttpResponseNotAllowed(['POST'])

    if 'file' not in request.FILES:
        return JsonResponse({'success': False, 'message': 'لم يتم رفع أي ملف'}, status=400)

    file = request.FILES['file']
    
    # Validate file extension
    if not file.name.endswith(('.csv', '.xlsx', '.xls')):
        return JsonResponse({'success': False, 'message': 'نوع الملف غير مدعوم'}, status=400)

    try:
        imported_count = 0
        errors = []

        if file.name.endswith('.csv'):
            # Handle CSV
            decoded_file = file.read().decode('utf-8-sig')
            csv_reader = csv.DictReader(io.StringIO(decoded_file))
            
            for row_num, row in enumerate(csv_reader, start=2):
                try:
                    Supplier.objects.create(
                        tenant=tenant,
                        name=row.get('name', '').strip() or row.get('الاسم', '').strip(),
                        phone=row.get('phone', '').strip() or row.get('الهاتف', '').strip() or None,
                        email=row.get('email', '').strip() or row.get('البريد', '').strip() or None,
                        city=row.get('city', '').strip() or row.get('المدينة', '').strip() or None,
                        address=row.get('address', '').strip() or row.get('العنوان', '').strip() or None,
                        opening_balance=float(row.get('opening_balance', 0) or row.get('الرصيد', 0) or 0),
                        credit_limit=float(row.get('credit_limit', 0) or row.get('حد_الائتمان', 0) or 0),
                        is_active=True
                    )
                    imported_count += 1
                except Exception as e:
                    errors.append(f'الصف {row_num}: {str(e)}')
        else:
            # Handle Excel - requires openpyxl
            try:
                import openpyxl
            except ImportError:
                return JsonResponse({
                    'success': False, 
                    'message': 'مكتبة openpyxl غير مثبتة. الرجاء تثبيتها أولاً'
                }, status=500)

            wb = openpyxl.load_workbook(file)
            ws = wb.active
            
            # Get headers from first row
            headers = [cell.value for cell in ws[1]]
            
            for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                try:
                    data = dict(zip(headers, row))
                    Supplier.objects.create(
                        tenant=tenant,
                        name=str(data.get('name', '') or data.get('الاسم', '')).strip(),
                        phone=str(data.get('phone', '') or data.get('الهاتف', '')).strip() or None,
                        email=str(data.get('email', '') or data.get('البريد', '')).strip() or None,
                        city=str(data.get('city', '') or data.get('المدينة', '')).strip() or None,
                        address=str(data.get('address', '') or data.get('العنوان', '')).strip() or None,
                        opening_balance=float(data.get('opening_balance', 0) or data.get('الرصيد', 0) or 0),
                        credit_limit=float(data.get('credit_limit', 0) or data.get('حد_الائتمان', 0) or 0),
                        is_active=True
                    )
                    imported_count += 1
                except Exception as e:
                    errors.append(f'الصف {row_num}: {str(e)}')

        message = f'تم استيراد {imported_count} مورد بنجاح'
        if errors:
            message += f'. حدثت {len(errors)} أخطاء'

        return JsonResponse({
            'success': True,
            'message': message,
            'imported': imported_count,
            'errors': errors[:10]  # Return first 10 errors only
        })

    except Exception as e:
        return JsonResponse({
            'success': False,
            'message': f'حدث خطأ أثناء الاستيراد: {str(e)}'
        }, status=500)


@login_required
def supplier_export_api(request):
    """Export suppliers to CSV file"""
    tenant = _ensure_tenant(request)
    if not tenant:
        return JsonResponse({'success': False, 'message': 'لا يوجد نشاط تجاري'}, status=400)

    # Create CSV response
    response = HttpResponse(content_type='text/csv; charset=utf-8-sig')
    response['Content-Disposition'] = 'attachment; filename="suppliers.csv"'
    
    # Add BOM for Excel UTF-8 support
    response.write('\ufeff')
    
    writer = csv.writer(response)
    
    # Write headers
    writer.writerow([
        'الاسم', 'الكود', 'الهاتف', 'البريد', 'المدينة', 'العنوان',
        'الرصيد الافتتاحي', 'حد الائتمان', 'الملاحظات', 'نشط'
    ])
    
    # Write data
    suppliers = Supplier.objects.for_tenant(tenant).order_by('name')
    for supplier in suppliers:
        writer.writerow([
            supplier.name,
            supplier.code,
            supplier.phone or '',
            supplier.email or '',
            supplier.city or '',
            supplier.address or '',
            supplier.opening_balance,
            supplier.credit_limit,
            supplier.notes or '',
            'نعم' if supplier.is_active else 'لا'
        ])
    
    return response


@login_required
def download_template(request):
    """Download CSV template for import"""
    response = HttpResponse(content_type='text/csv; charset=utf-8-sig')
    response['Content-Disposition'] = 'attachment; filename="suppliers_template.csv"'
    
    # Add BOM for Excel UTF-8 support
    response.write('\ufeff')
    
    writer = csv.writer(response)
    
    # Write headers
    writer.writerow([
        'الاسم', 'الهاتف', 'البريد', 'المدينة', 'العنوان',
        'الرصيد الافتتاحي', 'حد الائتمان'
    ])
    
    # Write example row
    writer.writerow([
        'شركة التوريدات المحدودة', '0512345678', 'supplier@example.com', 'الرياض', 'شارع الملك فهد',
        '0', '10000'
    ])
    
    return response
